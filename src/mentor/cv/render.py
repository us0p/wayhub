"""Rendering a CV (M9, D10, D81): the page, the PDF (WeasyPrint, from `cv/print.html`) and the
DOCX (python-docx) all come from the same `CVContent`. One ATS-friendly single-column layout:
no photo, no tables or columns, real text, standard section headings in the CV's language.

PDF and DOCX are rendered when downloaded (a CV is small; it takes well under a second), in a
worker thread so the event loop isn't blocked.
"""

import io
import re
import unicodedata
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from starlette.concurrency import run_in_threadpool

from mentor.cv.schemas import CVContent, ExperienceEntry
from mentor.web import templates

FONT = "Arial"  # Liberation Sans in the PDF: same metrics


@dataclass(frozen=True)
class Labels:
    summary: str
    experience: str
    skills: str
    education: str
    languages: str
    present: str
    months: tuple[str, ...]


LABELS = {
    "pt": Labels(
        "Resumo",
        "Experiência",
        "Habilidades",
        "Formação",
        "Idiomas",
        "atual",
        ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"),
    ),
    "en": Labels(
        "Summary",
        "Experience",
        "Skills",
        "Education",
        "Languages",
        "Present",
        ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
    ),
    "es": Labels(
        "Resumen",
        "Experiencia",
        "Habilidades",
        "Formación",
        "Idiomas",
        "actual",
        ("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sept", "oct", "nov", "dic"),
    ),
}


def labels(language: str) -> Labels:
    return LABELS.get(language.split("-")[0].lower(), LABELS["en"])


def _month(value: str | None, words: Labels) -> str | None:
    if not value:
        return None
    year, month = value.split("-")
    return f"{words.months[int(month) - 1]} {year}"


def period(entry: ExperienceEntry, words: Labels) -> str:
    start = _month(entry.start, words)
    end = words.present if entry.current else _month(entry.end, words)
    return " – ".join(x for x in (start, end) if x)  # noqa: RUF001 (typographic dash)


def contact_line(content: CVContent) -> list[str]:
    contact = content.contact
    return [x for x in (contact.location, contact.email, contact.phone) if x] + [
        link.url for link in contact.links
    ]


def filename(content: CVContent, job_title: str | None, extension: str) -> str:
    """'CV - Ana Souza - Desenvolvedora Python.pdf', ASCII-safe for the header."""
    parts = ["CV", content.contact.name or "", job_title or ""]
    text = " - ".join(p for p in parts if p)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"[^A-Za-z0-9 ._-]+", "", text).strip()[:120] or "CV"
    return f"{text}.{extension}"


def context(content: CVContent) -> dict[str, object]:
    words = labels(content.language)
    return {
        "cv": content,
        "words": words,
        "contact_line": contact_line(content),
        "periods": [period(e, words) for e in content.experiences],
    }


def _pdf(content: CVContent) -> bytes:
    from weasyprint import HTML  # heavy import, only when a PDF is asked for

    html = templates.get_template("cv/print.html").render(context(content))
    pdf: bytes = HTML(string=html).write_pdf()
    return pdf


async def pdf(content: CVContent) -> bytes:
    return await run_in_threadpool(_pdf, content)


def _heading(document: object, text: str) -> None:
    paragraph = document.add_paragraph()  # type: ignore[attr-defined]
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x2F, 0x2F, 0x2F)


def _docx(content: CVContent) -> bytes:
    words = labels(content.language)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = name.add_run(content.contact.name or "")
    run.bold = True
    run.font.size = Pt(18)
    if content.headline:
        document.add_paragraph(content.headline.text)
    if line := contact_line(content):
        document.add_paragraph(" | ".join(line))

    if content.summary:
        _heading(document, words.summary)
        document.add_paragraph(content.summary.text)
    if content.experiences:
        _heading(document, words.experience)
        for entry in content.experiences:
            title = document.add_paragraph()
            title.paragraph_format.space_before = Pt(6)
            title.add_run(entry.title.text).bold = True
            if entry.company:
                title.add_run(f" — {entry.company}")
            if when := period(entry, words):
                document.add_paragraph(when)
            for bullet in entry.bullets:
                document.add_paragraph(bullet.text, style="List Bullet")
    if content.skills:
        _heading(document, words.skills)
        document.add_paragraph(", ".join(skill.text for skill in content.skills))
    if content.education:
        _heading(document, words.education)
        for item in content.education:
            paragraph = document.add_paragraph()
            paragraph.add_run(item.text).bold = True
            if item.detail:
                paragraph.add_run(f" — {item.detail}")
    if content.languages:
        _heading(document, words.languages)
        for item in content.languages:
            document.add_paragraph(item.text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def docx(content: CVContent) -> bytes:
    return await run_in_threadpool(_docx, content)
